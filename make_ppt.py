# -*- coding: utf-8 -*-
"""
内部人员信息管理系统汇报PPT
严肃简洁风格，共7页
结构：封面 → 项目背景 → 系统概览 → 功能演示(现场) → 现存问题 → 未来展望 → 总结
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

OUT_DIR = r"e:\工作工作\text\person-info-system"
OUT_FILE = os.path.join(OUT_DIR, "内部人员信息管理系统汇报_最新.pptx")

# ===== 主题色：严肃深色系 =====
C_DARK   = RGBColor(0x1A, 0x2A, 0x3A)   # 深色背景
C_BLUE   = RGBColor(0x2B, 0x57, 0x8A)   # 主色蓝
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT  = RGBColor(0xF5, 0xF7, 0xFA)   # 浅灰背景
C_GRAY   = RGBColor(0x8C, 0x8C, 0x8C)   # 灰色文字
C_TEXT   = RGBColor(0x33, 0x33, 0x33)   # 正文
C_LINE   = RGBColor(0xD8, 0xDE, 0xE6)   # 分割线
C_ACCENT = RGBColor(0x3B, 0x7A, 0xB5)   # 点缀蓝

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_run(run, size=18, bold=False, color=C_TEXT, font_name="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn('a:ea'), {'typeface': font_name})
    rPr.append(ea)

def add_text(slide, left, top, width, height, text, size=18, bold=False,
             color=C_TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=bold, color=color)
    return tb

def add_rect(slide, left, top, width, height, fill=None, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.5)
    return shape

def add_page_header(slide, title, page_num):
    """每页顶部标题栏"""
    # 顶部蓝色细条
    add_rect(slide, 0, 0, prs.slide_width, Inches(0.06), fill=C_BLUE)
    # 标题
    add_text(slide, Inches(0.8), Inches(0.35), Inches(10), Inches(0.6),
             title, size=26, bold=True, color=C_BLUE)
    # 底部标题下划线
    add_rect(slide, Inches(0.8), Inches(1.05), Inches(1.2), Inches(0.04), fill=C_BLUE)
    # 页码
    add_text(slide, Inches(12.2), Inches(7.05), Inches(0.8), Inches(0.3),
             str(page_num), size=11, color=C_GRAY, align=PP_ALIGN.RIGHT)
    # 底部分割线
    add_rect(slide, Inches(0.8), Inches(7.1), Inches(11.7), Inches(0.02), fill=C_LINE)
    # 底部署名
    add_text(slide, Inches(0.8), Inches(7.15), Inches(8), Inches(0.25),
             "单位内部信息管理系统", size=9, color=C_GRAY)

def add_bullet(slide, left, top, width, height, items, size=16, color=C_TEXT):
    """简洁的无序号列表"""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        r0 = p.add_run()
        r0.text = "—  "
        set_run(r0, size=size, color=C_BLUE, bold=True)
        r1 = p.add_run()
        r1.text = item
        set_run(r1, size=size, color=color)
    return tb


# ======================================================================
# 第1页：封面
# ======================================================================
s1 = prs.slides.add_slide(BLANK)
# 深色背景
add_rect(s1, 0, 0, prs.slide_width, prs.slide_height, fill=C_DARK)
# 左侧蓝色竖条
add_rect(s1, 0, 0, Inches(0.15), prs.slide_height, fill=C_BLUE)
# 标题
add_text(s1, Inches(1.2), Inches(2.2), Inches(11), Inches(1.2),
         "内部人员信息管理系统", size=40, bold=True, color=C_WHITE)
# 副标题
add_text(s1, Inches(1.2), Inches(3.5), Inches(11), Inches(0.6),
         "项目建设情况汇报", size=22, color=RGBColor(0xA0, 0xB8, 0xD0))
# 分隔线
add_rect(s1, Inches(1.2), Inches(4.3), Inches(2.5), Inches(0.03), fill=C_BLUE)
# 底部信息
add_text(s1, Inches(1.2), Inches(5.8), Inches(11), Inches(0.4),
         "汇报人：信息中心", size=16, color=RGBColor(0x80, 0x98, 0xB0))
add_text(s1, Inches(1.2), Inches(6.3), Inches(11), Inches(0.4),
         "2026年8月", size=16, color=RGBColor(0x80, 0x98, 0xB0))


# ======================================================================
# 第2页：项目背景
# ======================================================================
s2 = prs.slides.add_slide(BLANK)
add_page_header(s2, "一、项目背景", 2)

add_text(s2, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.5),
         "传统办公模式下的主要痛点", size=18, bold=True, color=C_BLUE)

items_bg = [
    "人员信息分散管理：员工档案以纸质或Excel形式分散保存，查询效率低，更新不及时",
    "表单收集靠人工：工资条、考核表、信息采集等需逐个发邮件/微信群收集，汇总耗时",
    "审批流程纸质化：请假、报销、用章等申请需纸质流转，审批进度不透明，追溯困难",
    "数据重复录入：员工每次填表都需重复填写姓名、部门等基础信息，体验差、易出错",
    "信息安全隐患：邮件和微信群传输敏感信息，缺乏权限管控和操作审计",
]
add_bullet(s2, Inches(0.8), Inches(2.0), Inches(11.5), Inches(4.5), items_bg, size=16)

# 底部引出
add_rect(s2, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.6), fill=C_LIGHT)
add_text(s2, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.6),
         "目标：建设统一的内部人员信息管理系统，实现人员档案数字化、审批流程在线化、表单收集智能化",
         size=15, bold=True, color=C_BLUE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ======================================================================
# 第3页：系统概览
# ======================================================================
s3 = prs.slides.add_slide(BLANK)
add_page_header(s3, "二、系统概览", 3)

add_text(s3, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.5),
         "系统已实现的核心功能模块", size=18, bold=True, color=C_BLUE)

# 六大模块 - 简洁表格风格
modules = [
    ("人员档案管理", "全员电子档案，支持自定义字段、简历智能识别建档"),
    ("OA流程审批", "请假/报销/用章/出差四类流程，动态审批链路，退回修改"),
    ("智能表格", "上传文档自动识别字段，下发填写，原模板导出，一站式收集"),
    ("组织架构", "单位/部门/岗位/职级体系管理，审批人自动匹配"),
    ("通知公告", "在线发布通知，已读未读跟踪"),
    ("系统安全", "RBAC权限管控，密码策略，账号锁定，操作审计"),
]

for i, (title, desc) in enumerate(modules):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + col * Inches(5.9)
    y = Inches(2.0) + row * Inches(1.55)
    # 左侧蓝色竖条
    add_rect(s3, x, y, Inches(0.06), Inches(1.35), fill=C_BLUE)
    # 标题
    add_text(s3, x + Inches(0.25), y + Inches(0.1), Inches(5.4), Inches(0.4),
             title, size=17, bold=True, color=C_BLUE)
    # 描述
    add_text(s3, x + Inches(0.25), y + Inches(0.55), Inches(5.4), Inches(0.75),
             desc, size=14, color=C_TEXT)

# 技术架构
add_rect(s3, Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.04), fill=C_LINE)
add_text(s3, Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.3),
         "技术架构：Python FastAPI + Vue 3 + SQLite  |  单进程部署，支持内网服务器",
         size=13, color=C_GRAY)


# ======================================================================
# 第4页：功能演示
# ======================================================================
s4 = prs.slides.add_slide(BLANK)
add_page_header(s4, "三、功能演示", 4)

add_text(s4, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.5),
         "现场操作演示以下核心流程", size=18, bold=True, color=C_BLUE)

demos = [
    ("1", "人员档案", "查看人员档案 → 上传简历自动识别填充 → 人事批量导入建档"),
    ("2", "流程审批", "发起请假申请 → 部门领导审批 → 人事审批 → 查看流转进度"),
    ("3", "智能表格", "上传Word文档 → 智能识别字段 → 标注配置 → 下发填写 → 收集导出"),
]

for i, (num, title, desc) in enumerate(demos):
    y = Inches(2.2) + i * Inches(1.4)
    # 序号
    add_rect(s4, Inches(0.8), y, Inches(0.7), Inches(1.1), fill=C_BLUE)
    add_text(s4, Inches(0.8), y, Inches(0.7), Inches(1.1),
             num, size=28, bold=True, color=C_WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # 内容
    add_rect(s4, Inches(1.5), y, Inches(10.8), Inches(1.1), fill=C_LIGHT)
    add_text(s4, Inches(1.8), y + Inches(0.15), Inches(10.2), Inches(0.4),
             title, size=18, bold=True, color=C_BLUE)
    add_text(s4, Inches(1.8), y + Inches(0.6), Inches(10.2), Inches(0.4),
             desc, size=15, color=C_TEXT)

# 提示
add_text(s4, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4),
         "（现场打开系统网页进行操作演示）", size=14, color=C_GRAY, align=PP_ALIGN.CENTER)


# ======================================================================
# 第5页：现存问题
# ======================================================================
s5 = prs.slides.add_slide(BLANK)
add_page_header(s5, "四、现存问题", 5)

problems = [
    ("尚未上线运行", "系统目前在开发环境完成功能验证，尚未部署到内网服务器进行实际运行测试"),
    ("业务需求待深入了解", "刚入职不久，对各科室具体业务流程和实际使用需求了解还不充分，需要边用边摸索"),
    ("缺少统一消息中心", "通知公告、待办事项、审批提醒、填写任务等通知分散在各模块，缺乏统一的站内信消息入口"),
    ("缺少主动推送", "流程待办、填写任务通知仅依赖登录后查看，无邮件/系统弹窗等主动提醒方式"),
    ("PDF处理有限", "PDF文档已支持上传识别和标注，但预览渲染和回填精度仍需优化"),
]

for i, (title, desc) in enumerate(problems):
    y = Inches(1.5) + i * Inches(1.05)
    # 圆点
    add_rect(s5, Inches(0.9), y + Inches(0.08), Inches(0.12), Inches(0.12), fill=C_BLUE)
    # 标题
    add_text(s5, Inches(1.2), y, Inches(3.5), Inches(0.4),
             title, size=16, bold=True, color=C_BLUE)
    # 描述
    add_text(s5, Inches(4.8), y, Inches(7.5), Inches(0.4),
             desc, size=15, color=C_TEXT)

# 底部说明
add_rect(s5, Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.04), fill=C_LINE)
add_text(s5, Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.3),
         "当前策略：先部署上线、小范围试运行，收集实际使用反馈后再逐步完善和增加新功能",
         size=13, color=C_GRAY)


# ======================================================================
# 第6页：未来展望
# ======================================================================
s6 = prs.slides.add_slide(BLANK)
add_page_header(s6, "五、未来展望", 6)

# 近期（1-3个月）
add_text(s6, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.4),
         "近期计划（1-3个月）", size=17, bold=True, color=C_BLUE)
near_items = [
    "部署上线：部署到内网服务器，小范围试运行，收集实际使用反馈",
    "逐步完善：根据反馈修复问题，搭建统一消息中心、邮件提醒等基础体验",
    "深入了解需求：与各科室沟通，梳理实际业务流程和使用场景",
]
add_bullet(s6, Inches(0.8), Inches(1.8), Inches(11.5), Inches(1.6), near_items, size=15)

# 中期（3-6个月）
add_text(s6, Inches(0.8), Inches(3.6), Inches(11.5), Inches(0.4),
         "中期计划（3-6个月）", size=17, bold=True, color=C_BLUE)
mid_items = [
    "按需增加功能：根据试运行中发现的实际需求，逐步增加报表统计、模板市场等",
    "流程优化：与现有OA系统对接，支持自定义审批流程和电子签章",
    "数据迁移：根据使用规模，适时将数据库迁移至PostgreSQL/MySQL",
]
add_bullet(s6, Inches(0.8), Inches(4.1), Inches(11.5), Inches(1.6), mid_items, size=15)

# 长期
add_text(s6, Inches(0.8), Inches(5.9), Inches(11.5), Inches(0.4),
         "长期目标", size=17, bold=True, color=C_BLUE)
long_text = "持续迭代优化，逐步引入AI辅助、跨部门数据联动、国密加密审计，打造安全、智能、高效的人员信息管理平台"
add_text(s6, Inches(0.8), Inches(6.4), Inches(11.5), Inches(0.5),
         long_text, size=15, color=C_TEXT)


# ======================================================================
# 第7页：总结
# ======================================================================
s7 = prs.slides.add_slide(BLANK)
add_rect(s7, 0, 0, prs.slide_width, prs.slide_height, fill=C_DARK)
add_rect(s7, 0, 0, Inches(0.15), prs.slide_height, fill=C_BLUE)

add_text(s7, Inches(1.2), Inches(1.8), Inches(11), Inches(0.8),
         "总结", size=36, bold=True, color=C_WHITE)
add_rect(s7, Inches(1.2), Inches(2.7), Inches(1.5), Inches(0.04), fill=C_BLUE)

summary = [
    ("已完成", "人员档案、OA审批、智能表格、组织架构等核心模块开发完成，通过功能验证"),
    ("已保障", "权限校验、密码策略、路径防护等安全机制全面覆盖，满足内网部署要求"),
    ("下一步", "部署上线、小范围试运行，根据实际使用反馈逐步完善和增加新功能"),
]

for i, (tag, desc) in enumerate(summary):
    y = Inches(3.2) + i * Inches(1.0)
    add_text(s7, Inches(1.2), y, Inches(1.8), Inches(0.5),
             tag, size=18, bold=True, color=RGBColor(0x5B, 0x9B, 0xD5))
    add_text(s7, Inches(3.2), y, Inches(9), Inches(0.5),
             desc, size=17, color=C_WHITE)

add_rect(s7, Inches(5.5), Inches(6.5), Inches(2.3), Inches(0.03), fill=C_BLUE)
add_text(s7, Inches(1.2), Inches(6.65), Inches(11), Inches(0.5),
         "请领导批示，谢谢！", size=22, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)


# ===== 保存 =====
prs.save(OUT_FILE)
print(f"PPT saved to: {OUT_FILE}")
print(f"Size: {os.path.getsize(OUT_FILE)/1024:.1f} KB")
print(f"Total slides: {len(prs.slides)}")
