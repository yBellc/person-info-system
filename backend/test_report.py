"""模块3 报表流程测试：生成模板Excel→上传匹配→保存→生成报表"""
import os
import requests
from openpyxl import Workbook

BASE = "http://127.0.0.1:8000/api/v1"
UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "uploads")


def make_template_excel():
    """制造一个模拟上级下发的固定格式统计表模板"""
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    path = os.path.join(UPLOAD_DIR, "sample_official_template.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "人员统计表"
    # 表头（用各种常见表述，测试智能匹配）
    headers = ["序号", "姓名", "性别", "出生年月", "年龄", "身份证号",
               "所在单位", "职务", "文化程度", "参加工作时间", "联系电话"]
    ws.append(headers)
    # 留若干空行供填充
    for _ in range(20):
        ws.append([""] * len(headers))
    wb.save(path)
    wb.close()
    return path


def main():
    # 超管登录
    r = requests.post(f"{BASE}/auth/login", json={"username": "admin", "password": "admin123"})
    assert r.status_code == 200, r.text
    token = r.json()["access_token"]
    H = {"Authorization": f"Bearer {token}"}

    # 1. 制造模板文件
    tpl_path = make_template_excel()
    print(f"✓ 生成测试模板: {tpl_path}")

    # 2. 上传并自动匹配
    with open(tpl_path, "rb") as f:
        r = requests.post(
            f"{BASE}/reports/upload-and-match?header_row=1",
            headers=H,
            files={"file": ("sample_official_template.xlsx", f,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert r.status_code == 200, r.text
    data = r.json()
    print(f"\n✓ 上传并自动匹配结果:")
    print(f"  表头: {data['headers']}")
    print(f"  匹配:")
    for m in data["mappings"]:
        flag = "🟢" if m["confirmed"] else ("🟡" if m["system_field_key"] else "🔴")
        print(f"    {flag} '{m['template_name']}' -> '{m['system_field_key']}' (置信度={m['confidence']})")

    # 3. 保存模板（确认映射后）
    mappings_payload = [
        {"col_index": m["col_index"], "template_name": m["template_name"],
         "system_field_key": m["system_field_key"] or None,
         "confidence": m["confidence"], "confirmed": True}
        for m in data["mappings"] if m["system_field_key"]
    ]
    r = requests.post(f"{BASE}/reports/templates", headers=H, json={
        "name": "2026年度人员情况统计表",
        "header_row": 1, "data_start_row": 2, "data_start_col": 1,
        "is_public": True,
        "mappings": mappings_payload,
    })
    assert r.status_code == 200, r.text
    template_id = r.json()["id"]
    print(f"\n✓ 保存模板成功 id={template_id}")

    # 4. 上传模板文件绑定
    with open(tpl_path, "rb") as f:
        r = requests.post(
            f"{BASE}/reports/templates/{template_id}/upload-file",
            headers=H,
            files={"file": ("tpl.xlsx", f,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert r.status_code == 200, r.text
    print(f"✓ 模板文件已绑定")

    # 5. 生成报表
    r = requests.post(f"{BASE}/reports/generate", headers=H, json={
        "template_id": template_id, "unit_id": 1,
    })
    assert r.status_code == 200, r.text
    result = r.json()
    print(f"\n✓ 生成报表结果: {result}")

    # 6. 下载报表
    if "output_file" in result:
        r = requests.get(f"{BASE}/reports/download/{result['output_file']}", headers=H)
        assert r.status_code == 200, r.text
        out_path = os.path.join(UPLOAD_DIR, "downloaded_report.xlsx")
        with open(out_path, "wb") as f:
            f.write(r.content)
        print(f"✓ 下载报表到: {out_path}")

        # 验证内容
        from openpyxl import load_workbook
        wb = load_workbook(out_path)
        ws = wb.active
        print(f"\n  报表内容预览（前5行）:")
        for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            print(f"    {row}")
        wb.close()

    print("\n===== 模块3 报表测试通过 =====")


if __name__ == "__main__":
    main()
